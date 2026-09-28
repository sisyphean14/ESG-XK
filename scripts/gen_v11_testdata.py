"""生成 ESG-XK V1.1.0 测试集：Office 内嵌图片 OCR + 可靠性（超时 / 卡死 / 自动跳过）。

输出到 D:\\ESG-XK\\ceshi：
  office/        DOCX / PPTX / XLSX 内嵌图片语料（只读，基线校验）
  office_versions/  图片变化的 v1 / v2 模板
  office_runctl/ 多图片 PPTX（安全停止 / 急停）
  reliability/   A–E 故障批次、加密 PDF、损坏 Office、空文件等
  V11_BASELINE.json  上述文件的 size / mtime / SHA-256 基线
"""
from __future__ import annotations

import copy
import hashlib
import io
import json
import random
import shutil
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont

ROOT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(r"D:\ESG-XK\ceshi")
FONT = r"C:\Windows\Fonts\msyh.ttc"
random.seed(20261001)

UNIQUE_DOCX = "ESG-XK-OCR-UNIQUE-TEST"
UNIQUE_PPTX = "ESG-XK-PPTX-OCR-UNIQUE"
UNIQUE_XLSX = "ESG-XK-XLSX-OCR-UNIQUE"


def text_img(lines: list[str], w: int = 1400, size: int = 44, bg=(255, 255, 255), fg=(20, 20, 20),
             border: bool = True) -> Image.Image:
    f = ImageFont.truetype(FONT, size)
    h = 70 + int(size * 1.8) * len(lines)
    img = Image.new("RGB", (w, h), bg)
    d = ImageDraw.Draw(img)
    y = 35
    for ln in lines:
        d.text((40, y), ln, fill=fg, font=f)
        y += int(size * 1.8)
    if border:
        d.rectangle([2, 2, w - 3, h - 3], outline=(160, 170, 180), width=3)
    return img


def png_bytes(img: Image.Image, fmt: str = "PNG", **kw) -> bytes:
    b = io.BytesIO()
    img.save(b, fmt, **kw)
    return b.getvalue()


def stream(data: bytes) -> io.BytesIO:
    return io.BytesIO(data)


IMG = {
    "zh": png_bytes(text_img(["职业健康安全管理制度（截图）", "所有员工每年接受不少于8小时安全培训"])),
    "en": png_bytes(text_img(["Scope 3 emissions from purchased goods", "reached 45,600 tCO2e in 2025"])),
    "id": png_bytes(text_img(["Target pengurangan emisi 30 persen", "pada tahun 2030 untuk seluruh pabrik"])),
    "unique_docx": png_bytes(text_img([UNIQUE_DOCX, "职业健康专项检查通知（仅存在于图片中）"])),
    "unique_pptx": png_bytes(text_img([UNIQUE_PPTX, "Safety walk findings closed within 7 days"])),
    "unique_xlsx": png_bytes(text_img([UNIQUE_XLSX, "Limbah B3 disimpan maksimal 90 hari"])),
    "dup": png_bytes(text_img(["重复图片：危险化学品仓库巡检记录", "DUP-IMG-2025 每周一次"])),
    "table": png_bytes(text_img(["表格内图片：污水排放口编号 WS-07"], w=1000, size=38)),
    "anchor": png_bytes(text_img(["浮动图片：碳排放因子采用2025年电网因子"], w=1300, size=40)),
    "vml": png_bytes(text_img(["VML图片：年度能源审计报告已完成"], w=1200, size=40)),
    "flow": png_bytes(text_img(["流程图截图：隐患上报 → 评估 → 整改 → 验证"], w=1500, size=40)),
    "v1": png_bytes(text_img(["版本图片 V1：年度节水目标 100 万吨"], w=1200, size=42)),
    "v2": png_bytes(text_img(["版本图片 V2：年度节水目标 200 万吨"], w=1200, size=42)),
}
# 旋转 90° 的截图（像素内容侧向）
IMG["rotated"] = png_bytes(text_img(["旋转截图：应急演练每半年开展一次"], w=1300, size=46).rotate(90, expand=True))
# 低清晰度
low = text_img(["低清晰度：排污许可证 GD-2025-0917"], w=1200, size=40)
low = low.resize((300, 38)).resize((900, 114)).filter(ImageFilter.GaussianBlur(1.6))
IMG["lowres"] = png_bytes(low, "JPEG", quality=15)
# 过滤对象
IMG["tiny"] = png_bytes(Image.new("RGB", (20, 20), (200, 30, 30)))
icon = Image.new("RGB", (64, 64), (10, 124, 102))
ImageDraw.Draw(icon).ellipse([10, 10, 54, 54], fill=(255, 255, 255))
IMG["icon"] = png_bytes(icon)
IMG["transparent"] = png_bytes(Image.new("RGBA", (400, 100), (0, 0, 0, 0)))
IMG["pixel"] = png_bytes(Image.new("RGB", (1, 1), (0, 0, 0)))
IMG["blank"] = png_bytes(Image.new("RGB", (800, 200), (255, 255, 255)))
# 格式覆盖
IMG["jpg"] = png_bytes(text_img(["JPEG 截图：绿色采购比例 62%"], w=1100, size=40), "JPEG", quality=90)
IMG["bmp"] = png_bytes(text_img(["BMP 截图：危废台账月度复核"], w=1100, size=40), "BMP")
IMG["tiff"] = png_bytes(text_img(["TIFF 截图：土壤监测点 12 个"], w=1100, size=40), "TIFF")


# ---------------------------------------------------------------- DOCX
def docx_add_anchor_picture(doc, run, data: bytes, width_cm: float):
    """先插入 inline 图片，再把 wp:inline 转换为 wp:anchor（浮动图片）。"""
    from docx.oxml.ns import qn
    from docx.shared import Cm
    run.add_picture(stream(data), width=Cm(width_cm))
    inline = run._r.find(qn("w:drawing")).find(qn("wp:inline"))
    from lxml import etree
    anchor = etree.SubElement(inline.getparent(), qn("wp:anchor"), {
        "distT": "0", "distB": "0", "distL": "114300", "distR": "114300", "simplePos": "0", "relativeHeight": "1",
        "behindDoc": "0", "locked": "0", "layoutInCell": "1", "allowOverlap": "1"})
    etree.SubElement(anchor, qn("wp:simplePos"), {"x": "0", "y": "0"})
    ph = etree.SubElement(anchor, qn("wp:positionH"), {"relativeFrom": "column"})
    etree.SubElement(ph, qn("wp:posOffset")).text = "0"
    pv = etree.SubElement(anchor, qn("wp:positionV"), {"relativeFrom": "paragraph"})
    etree.SubElement(pv, qn("wp:posOffset")).text = "0"
    for child in list(inline):
        if child.tag == qn("wp:extent"):
            anchor.append(copy.deepcopy(child))
            etree.SubElement(anchor, qn("wp:wrapSquare"), {"wrapText": "bothSides"})
        elif child.tag in (qn("wp:docPr"), qn("wp:cNvGraphicFramePr"), qn("a:graphic")):
            anchor.append(copy.deepcopy(child))
    inline.getparent().remove(inline)


def docx_add_vml_picture(doc, paragraph, rid: str):
    from docx.oxml import parse_xml
    xml = ('<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
           'xmlns:v="urn:schemas-microsoft-com:vml" '
           'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
           '<w:pict><v:shape id="vmlpic1" style="width:300pt;height:40pt" type="#_x0000_t75">'
           f'<v:imagedata r:id="{rid}" o:title="vml" xmlns:o="urn:schemas-microsoft-com:office:office"/>'
           '</v:shape></w:pict></w:r>')
    paragraph._p.append(parse_xml(xml))


def build_docx_main(path: Path) -> None:
    import docx
    from docx.shared import Cm
    d = docx.Document()
    d.add_heading("职业安全与环境管理手册（Office 图片 OCR 测试）", 0)
    d.add_heading("1 概述", 1)
    d.add_paragraph("本手册正文介绍安全与环境管理的总体要求。")          # 正文不含“职业健康”与唯一标记
    d.add_paragraph("下图为中文制度截图：")
    d.add_paragraph().add_run().add_picture(stream(IMG["zh"]), width=Cm(14))
    d.add_paragraph("以下为英文截图与印尼语截图。")
    d.add_paragraph().add_run().add_picture(stream(IMG["en"]), width=Cm(14))
    d.add_paragraph().add_run().add_picture(stream(IMG["id"]), width=Cm(14))
    d.add_heading("2 专项检查", 1)
    d.add_paragraph("专项检查通知见下方附图。")
    d.add_paragraph().add_run().add_picture(stream(IMG["unique_docx"]), width=Cm(14))
    d.add_heading("3 重复图片", 1)
    for i in range(3):
        d.add_paragraph(f"第{i + 1}处引用巡检记录：")
        d.add_paragraph().add_run().add_picture(stream(IMG["dup"]), width=Cm(12))
    d.add_heading("4 过滤对象", 1)
    p = d.add_paragraph("图标 ")
    p.add_run().add_picture(stream(IMG["icon"]), width=Cm(0.5))
    p.add_run(" 小图 ").add_picture(stream(IMG["tiny"]), width=Cm(0.4))
    d.add_paragraph().add_run().add_picture(stream(IMG["transparent"]), width=Cm(8))
    d.add_paragraph().add_run().add_picture(stream(IMG["pixel"]), width=Cm(0.2))
    d.add_paragraph().add_run().add_picture(stream(IMG["blank"]), width=Cm(8))
    d.add_heading("5 表格与图片", 1)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "监测项目"
    t.cell(0, 1).text = "现场照片"
    t.cell(1, 0).text = "污水排放口"
    t.cell(1, 1).paragraphs[0].add_run().add_picture(stream(IMG["table"]), width=Cm(6))
    d.add_heading("6 低清晰度与浮动图片", 1)
    d.add_paragraph("低清晰度扫描件：")
    d.add_paragraph().add_run().add_picture(stream(IMG["lowres"]), width=Cm(10))
    p = d.add_paragraph("浮动图片所在段落。")
    docx_add_anchor_picture(d, p.add_run(), IMG["anchor"], 13)
    d.add_heading("7 VML 与多格式", 1)
    pic_para = d.add_paragraph()
    pic_para.add_run().add_picture(stream(IMG["jpg"]), width=Cm(10))
    rid = d.part.relate_to(d.part.package.get_or_add_image_part(stream(IMG["vml"])).partname and
                           d.part.package.get_or_add_image_part(stream(IMG["vml"])), docx.opc.constants.RELATIONSHIP_TYPE.IMAGE)
    docx_add_vml_picture(d, d.add_paragraph("VML 图片："), rid)
    d.add_paragraph().add_run().add_picture(stream(IMG["bmp"]), width=Cm(10))
    d.add_paragraph().add_run().add_picture(stream(IMG["tiff"]), width=Cm(10))
    d.add_paragraph("手册结束。")
    d.save(str(path))


def build_docx_simple(path: Path, title: str, body: str, images: list[str]) -> None:
    import docx
    from docx.shared import Cm
    d = docx.Document()
    d.add_heading(title, 0)
    d.add_paragraph(body)
    for k in images:
        d.add_paragraph(f"附图（{k}）：")
        d.add_paragraph().add_run().add_picture(stream(IMG[k]), width=Cm(12))
    d.save(str(path))


# ---------------------------------------------------------------- PPTX
def build_pptx_main(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Cm, Pt
    prs = Presentation()
    blank = prs.slide_layouts[5]
    s = prs.slides.add_slide(blank)
    s.shapes.title.text = "培训材料：安全与排放"
    s.shapes.add_picture(stream(IMG["zh"]), Cm(1), Cm(4), width=Cm(11))
    s.shapes.add_picture(stream(IMG["en"]), Cm(13), Cm(4), width=Cm(11))
    tb = s.shapes.add_textbox(Cm(1), Cm(12), Cm(20), Cm(2))
    tb.text_frame.text = "文字 Shape 与图片混排：本页介绍安全培训与范围三排放。"
    s = prs.slides.add_slide(blank)
    s.shapes.title.text = "Kebijakan dan diagram"
    s.shapes.add_picture(stream(IMG["id"]), Cm(1), Cm(4), width=Cm(11))
    s.shapes.add_picture(stream(IMG["rotated"]), Cm(14), Cm(3), height=Cm(13))
    s.shapes.add_picture(stream(IMG["icon"]), Cm(23), Cm(1), width=Cm(0.8))
    s = prs.slides.add_slide(blank)
    s.shapes.title.text = "通知截图"
    s.shapes.add_picture(stream(IMG["unique_pptx"]), Cm(1), Cm(4), width=Cm(16))
    grp = s.shapes.add_group_shape()
    grp.shapes.add_picture(stream(IMG["flow"]), Cm(1), Cm(10), width=Cm(16))
    rot = s.shapes.add_picture(stream(IMG["dup"]), Cm(18), Cm(4), width=Cm(7))
    rot.rotation = 180.0
    tb = s.shapes.add_textbox(Cm(1), Cm(15), Cm(20), Cm(1.5))
    tb.text_frame.text = "本页包含通知截图与流程图截图。"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    prs.save(str(path))


def build_pptx_many(path: Path, n_slides: int = 12) -> None:
    from pptx import Presentation
    from pptx.util import Cm
    prs = Presentation()
    for i in range(1, n_slides + 1):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = f"运行控制测试 {i}"
        for j in (1, 2):
            n = (i - 1) * 2 + j
            img = png_bytes(text_img([f"安全停止测试图片 {n:02d}", f"巡检记录编号 RC-{1000 + n}"], w=1100, size=40))
            s.shapes.add_picture(stream(img), Cm(1 + (j - 1) * 12), Cm(5), width=Cm(11))
    prs.save(str(path))


# ---------------------------------------------------------------- XLSX
def build_xlsx(path: Path) -> None:
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "排放台账"
    ws.append(["年份", "范围一", "备注"])
    ws.append([2025, 12345, "见右侧截图"])
    im = XLImage(stream(IMG["zh"]))
    im.width, im.height = 560, 110
    ws.add_image(im, "B4")
    ws2 = wb.create_sheet("Water")
    ws2.append(["Year", "Withdrawal"])
    ws2.append([2025, 456000])
    im2 = XLImage(stream(IMG["en"]))
    im2.width, im2.height = 560, 110
    ws2.add_image(im2, "D5")
    im3 = XLImage(stream(IMG["unique_xlsx"]))
    im3.width, im3.height = 560, 110
    ws2.add_image(im3, "A12")
    ws3 = wb.create_sheet("Icon")
    im4 = XLImage(stream(IMG["icon"]))
    im4.width, im4.height = 16, 16
    ws3.add_image(im4, "C2")
    wb.save(str(path))


def build_encrypted_pdf(path: Path) -> None:
    from reportlab.lib.pdfencrypt import StandardEncryption
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path), encrypt=StandardEncryption("user-secret", "owner-secret", canPrint=0))
    c.drawString(72, 720, "Encrypted ESG document")
    c.save()


def main() -> None:
    for sub in ("office", "office_versions", "office_runctl", "reliability"):
        if (ROOT / sub).exists():
            shutil.rmtree(ROOT / sub)
        (ROOT / sub).mkdir(parents=True)
    off = ROOT / "office"
    build_docx_main(off / "职业安全环境手册_图片OCR.docx")
    build_docx_simple(off / "重复图片_文档A.docx", "巡检文档A", "文档A正文。", ["dup"])
    build_docx_simple(off / "重复图片_文档B.docx", "巡检文档B", "文档B正文。", ["dup"])
    build_pptx_main(off / "安全培训材料_图片OCR.pptx")
    build_xlsx(off / "排放台账_图片OCR.xlsx")
    ver = ROOT / "office_versions"
    import docx  # noqa: F401
    build_docx_simple(ver / "v1_节水方案.docx", "节水方案", "节水方案正文（版本测试）。", ["zh", "v1"])
    build_docx_simple(ver / "v2_节水方案.docx", "节水方案", "节水方案正文（版本测试）。", ["zh", "v2"])
    build_pptx_many(ROOT / "office_runctl" / "多图片培训_24张.pptx")
    rel = ROOT / "reliability"
    (rel / "A_正常文档.txt").write_text("A 正常文档：可靠性测试批次 A，年度安全培训完成率 99%。\n", encoding="utf-8")
    build_docx_simple(rel / "B_FAULT_PARSE_HANG.docx", "B 文档", "B 解析卡死测试文档。", ["en"])
    (rel / "C_正常文档.md").write_text("# C 正常文档\n\n可靠性测试批次 C：危废处置合规率 100%。\n", encoding="utf-8")
    text_img(["D OCR 卡死测试图片"], w=900, size=40).save(rel / "D_FAULT_OCR_HANG.png")
    (rel / "E_正常文档.txt").write_text("E 正常文档：可靠性测试批次 E，取水量同比下降 6%。\n", encoding="utf-8")
    build_encrypted_pdf(rel / "加密文档.pdf")
    (rel / "损坏Office.docx").write_bytes(b"PK\x03\x04" + bytes(random.getrandbits(8) for _ in range(3000)))
    (rel / "零字节.txt").write_bytes(b"")
    text_img(["F OCR 进程崩溃测试"], w=900, size=40).save(rel / "F_FAULT_OCR_CRASH.png")
    (rel / "G_FAULT_EMBED_HANG.txt").write_text("G 向量化卡死测试：本文件用于验证 Embedding 超时。\n" * 3, encoding="utf-8")
    (rel / "H_FAULT_WORKER_HANG.txt").write_text("H Worker 卡死测试：用于验证 Watchdog。\n", encoding="utf-8")
    # 多页扫描 PDF：单页卡死测试（第 2 页）
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(rel / "P_FAULT_PAGE_HANG_3页.pdf"), pagesize=A4)
    for i in range(1, 4):
        buf = io.BytesIO()
        text_img([f"单页超时测试 第{i}页", f"页面记录编号 PG-{i:03d}"], w=1400, size=48).save(buf, "PNG")
        buf.seek(0)
        c.drawImage(ImageReader(buf), 40, 500, width=515, height=200)
        c.showPage()
    c.save()
    baseline = {}
    for sub in ("office", "office_versions", "office_runctl", "reliability"):
        for p in sorted((ROOT / sub).rglob("*")):
            if p.is_file():
                st = p.stat()
                baseline[str(p)] = {"size": st.st_size, "mtime": st.st_mtime,
                                    "sha256": hashlib.sha256(p.read_bytes()).hexdigest()}
    (ROOT / "V11_BASELINE.json").write_text(json.dumps(baseline, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"generated {len(baseline)} files")


if __name__ == "__main__":
    main()
