"""PPTX 解析：每张幻灯片的标题/正文、表格、备注。"""
from __future__ import annotations

import io
import zipfile

from esgxk.core.errors import ErrorCode, EsgError
from esgxk.parsers.base import ParseResult, SourceUnit


def _shape_texts(shape, out: list[str], tables: list[list[str]]) -> None:
    if getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):  # GROUP
        for s in shape.shapes:
            _shape_texts(s, out, tables)
        return
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = []
        for row in shape.table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            tables.append(rows)
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        t = "\n".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
        if t:
            out.append(t)


def parse_pptx(data: bytes) -> ParseResult:
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise EsgError(ErrorCode.PARSE_CORRUPTED, f"PPTX 结构无效：{exc}") from exc
    except Exception as exc:
        raise EsgError(ErrorCode.PARSE_CORRUPTED, f"PPTX 打开失败：{type(exc).__name__}: {exc}") from exc
    result = ParseResult()
    for idx, slide in enumerate(prs.slides, start=1):
        title = ""
        try:
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                title = slide.shapes.title.text_frame.text.strip()
        except Exception:
            title = ""
        texts: list[str] = []
        tables: list[list[str]] = []
        for shape in slide.shapes:
            _shape_texts(shape, texts, tables)
        body = [t for t in texts if t != title]
        if title or body:
            text = (title + "\n" if title else "") + "\n".join(body)
            loc = {"type": "pptx", "slide": idx}
            if title:
                loc["heading"] = title
            result.units.append(SourceUnit(
                unit_type="slide_text", text=text.strip(), slide=idx, heading_path=title or None, locator=loc,
            ))
        for ti, rows in enumerate(tables, start=1):
            loc = {"type": "pptx", "slide": idx, "table": ti}
            if title:
                loc["heading"] = title
            result.units.append(SourceUnit(
                unit_type="table", text="\n".join(rows), slide=idx, heading_path=title or None, locator=loc,
            ))
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    result.units.append(SourceUnit(
                        unit_type="slide_notes", text=notes, slide=idx, heading_path=title or None,
                        locator={"type": "pptx", "slide": idx, "part": "notes"},
                    ))
        except Exception:
            result.warnings.append(f"第{idx}张幻灯片备注读取失败")
    result.page_count = len(prs.slides)
    return result
